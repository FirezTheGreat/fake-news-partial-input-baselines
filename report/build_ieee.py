"""Converts the final report (report/WSMA_Report.docx, hand-edited) into IEEE conference format:
report/WSMA_Report_IEEE.docx. The text is kept as written; only the layout, numbering, captions,
citations and references change. Export to PDF from Word afterwards.

    .venv\\Scripts\\python report\\build_ieee.py

IEEE A4 layout: title and author block across the page, two columns below (Times New Roman 10 pt),
sections I, II ... with subsections A, B ..., "Fig. n." captions below figures, "TABLE I" captions
above tables, numbered citations [n] in order of first use. Charts are redrawn without their in-chart
titles (the captions replace them): five at column width; the wide figures and tables span both
columns.
"""

import copy
import importlib
import io
import re
import sys
import tempfile
from pathlib import Path

import docx
import matplotlib
import matplotlib.figure
import matplotlib.pyplot as plt
from PIL import Image
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "report" / "WSMA_Report.docx"
OUT = ROOT / "report" / "WSMA_Report_IEEE.docx"
sys.path.insert(0, str(ROOT / "scripts"))

FONT = "Times New Roman"
COL_W = Inches(3.48)    # one column on IEEE A4 (210 mm - 2 x 14.32 mm margins - 4.22 mm gap) / 2
FULL_W = Inches(7.0)    # both columns

TITLE = "Partial-Input Baselines Reveal Construction Artifacts in Fake News Benchmarks"
AUTHOR = ["Abhishek Ghosh",
          "Dept. of Computer Science and Engineering (Data Science)",
          "Manipal Institute of Technology, Bengaluru, India",
          "Reg. No. 245816186"]

# Images in document order: (source file to use, full width?, column-size redraw function or None)
IMAGES = [
    ("report/architecture.png", True, None),
    ("figures/fig1_accuracy.png", False, "fig1_accuracy"),
    ("figures/fig2_leakage_tokens.png", True, "fig2_tokens"),
    ("figures/fig3_confusion.png", True, "fig3_confusion"),
    ("figures/fig7_cross_corpus.png", False, "fig7_cross_corpus"),
    ("figures/fig5_topic_matched.png", False, "fig5_topic_matched"),
    ("figures/fig6_svm_check.png", False, "fig6_svm_check"),
    ("figures/fig4_degradation.png", False, "fig4_degradation"),
]
REDRAWN = [redraw for _, _, redraw in IMAGES if redraw]
# (width, height) factors for each redrawn chart. Fonts keep their size, so a smaller canvas
# gives relatively larger text once the chart is scaled to the column.
CHART_SCALE = {"fig1_accuracy": (0.58, 0.72), "fig2_tokens": (1.0, 0.6), "fig3_confusion": (1.0, 1.0),
               "fig4_degradation": (0.85, 0.72), "fig5_topic_matched": (0.6, 0.72),
               "fig6_svm_check": (0.62, 0.72), "fig7_cross_corpus": (0.6, 0.72)}
CHART_FILE = {"fig2_tokens": "fig2_leakage_tokens.png"}
current = [None]
FULL_WIDTH_TABLES = {0, 1}   # Table I (related work) and Table II (all results)

# Citations in order of first use -> number; (pattern, replacement) applied to the text.
CITATIONS = [
    (r"Shu et al\. \(2017\)", "Shu et al. [1]"),
    (r"Ahmed, Traore and Saad \(2017\)", "Ahmed, Traore and Saad [2]"),
    (r"Wang \(2017\)", "Wang [3]"),
    (r"McIntire's \(2017\)", "McIntire's [4]"),
    (r"Gururangan et al\. \(2018\)", "Gururangan et al. [5]"),
    (r"Poliak et al\. \(2018\)", "Poliak et al. [6]"),
    (r"McCoy, Pavlick and Linzen \(2019\)", "McCoy, Pavlick and Linzen [7]"),
    (r"McCoy et al\. \(2019\)", "McCoy et al. [7]"),
    (r"Niven (?:and|&) Kao \(2019\)", "Niven and Kao [8]"),
    (r"Schuster et al\. \(2019\)", "Schuster et al. [9]"),
    (r"Geirhos et al\. \(2020\)", "Geirhos et al. [10]"),
    (r"\(Zhou et al\., 2021\)", "[11]"),
    (r"Zhou et al\. \(2021\)", "Zhou et al. [11]"),
    (r"\(Verma, 2026\)", "[12]"),
    (r"Verma's \(2026\)", "Verma's [12]"),
    (r"Verma \(2026\)", "Verma [12]"),
    (r"Kaushik, Hovy and Lipton \(2020\)", "Kaushik, Hovy and Lipton [13]"),
    (r"Kaushik et al\. \(2020\)", "Kaushik et al. [13]"),
    (r"Gardner et al\. \(2020\)", "Gardner et al. [14]"),
]

# IEEE reference list: segments of (text, italic).
REFERENCES = [
    [("K. Shu, A. Sliva, S. Wang, J. Tang, and H. Liu, “Fake news detection on social media: A data "
      "mining perspective,” ", False), ("ACM SIGKDD Explorations Newsl.", True),
     (", vol. 19, no. 1, pp. 22–36, 2017, doi: 10.1145/3137597.3137600.", False)],
    [("H. Ahmed, I. Traore, and S. Saad, “Detection of online fake news using n-gram analysis and "
      "machine learning techniques,” in ", False),
     ("Intelligent, Secure, and Dependable Systems in Distributed and Cloud Environments (ISDDC 2017)", True),
     (", ser. Lecture Notes in Computer Science, vol. 10618. Cham, Switzerland: Springer, 2017, "
      "pp. 127–138.", False)],
    [("W. Y. Wang, “‘Liar, liar pants on fire’: A new benchmark dataset for fake news detection,” in ", False),
     ("Proc. 55th Annu. Meeting Assoc. Comput. Linguistics (Vol. 2: Short Papers)", True),
     (", 2017, pp. 422–426, doi: 10.18653/v1/P17-2067.", False)],
    [("G. McIntire, “Fake or real news dataset,” GitHub, 2017. [Online]. Available: "
      "https://github.com/joolsa/fake_real_news_dataset", False)],
    [("S. Gururangan, S. Swayamdipta, O. Levy, R. Schwartz, S. R. Bowman, and N. A. Smith, “Annotation "
      "artifacts in natural language inference data,” in ", False),
     ("Proc. NAACL-HLT 2018 (Vol. 2: Short Papers)", True),
     (", 2018, pp. 107–112, doi: 10.18653/v1/N18-2017.", False)],
    [("A. Poliak, J. Naradowsky, A. Haldar, R. Rudinger, and B. Van Durme, “Hypothesis only baselines in "
      "natural language inference,” in ", False),
     ("Proc. 7th Joint Conf. Lexical and Computational Semantics (*SEM)", True),
     (", 2018, pp. 180–191, doi: 10.18653/v1/S18-2023.", False)],
    [("R. T. McCoy, E. Pavlick, and T. Linzen, “Right for the wrong reasons: Diagnosing syntactic "
      "heuristics in natural language inference,” in ", False),
     ("Proc. 57th Annu. Meeting Assoc. Comput. Linguistics", True),
     (", 2019, pp. 3428–3448, doi: 10.18653/v1/P19-1334.", False)],
    [("T. Niven and H.-Y. Kao, “Probing neural network comprehension of natural language arguments,” in ", False),
     ("Proc. 57th Annu. Meeting Assoc. Comput. Linguistics", True),
     (", 2019, pp. 4658–4664, doi: 10.18653/v1/P19-1459.", False)],
    [("T. Schuster, D. Shah, Y. J. S. Yeo, D. R. Filizzola Ortiz, E. Santus, and R. Barzilay, “Towards "
      "debiasing fact verification models,” in ", False),
     ("Proc. EMNLP-IJCNLP 2019", True), (", 2019, pp. 3419–3425, doi: 10.18653/v1/D19-1341.", False)],
    [("R. Geirhos, J.-H. Jacobsen, C. Michaelis, R. Zemel, W. Brendel, M. Bethge, and F. A. Wichmann, "
      "“Shortcut learning in deep neural networks,” ", False), ("Nature Mach. Intell.", True),
     (", vol. 2, no. 11, pp. 665–673, 2020, doi: 10.1038/s42256-020-00257-z.", False)],
    [("X. Zhou, H. Elfardy, C. Christodoulopoulos, T. Butler, and M. Bansal, “Hidden biases in unreliable "
      "news detection datasets,” in ", False),
     ("Proc. 16th Conf. Eur. Chapter Assoc. Comput. Linguistics (EACL)", True),
     (", 2021, pp. 2482–2492, doi: 10.18653/v1/2021.eacl-main.211.", False)],
    [("Y. Verma, “What does 99% accuracy measure? A reproducible audit of shortcut learning in a widely "
      "used fake news corpus,” 2026, ", False), ("arXiv:2609.25006", False),
     (", doi: 10.48550/arXiv.2609.25006.", False)],
    [("D. Kaushik, E. Hovy, and Z. C. Lipton, “Learning the difference that makes a difference with "
      "counterfactually-augmented data,” in ", False),
     ("Proc. Int. Conf. Learning Representations (ICLR)", True), (", 2020, arXiv:1909.12434.", False)],
    [("M. Gardner ", False), ("et al.", True), (", “Evaluating models’ local decision boundaries via contrast "
      "sets,” in ", False), ("Findings Assoc. Comput. Linguistics: EMNLP 2020", True),
     (", 2020, pp. 1307–1323, doi: 10.18653/v1/2020.findings-emnlp.117.", False)],
]

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
LETTER = " ABCDEFGHIJ"


# ---------------------------------------------------------------- helpers

def replace_in_paragraph(p, pattern, repl):
    """Regex replace across the runs of a paragraph, keeping each run's formatting."""
    runs = p.runs
    text = "".join(r.text for r in runs)
    matches = list(re.finditer(pattern, text))
    for m in reversed(matches):
        new = m.expand(repl) if isinstance(repl, str) else repl(m)
        s, e, pos, first = m.start(), m.end(), 0, True
        for r in runs:
            rs, rend = pos, pos + len(r.text)
            pos = rend
            if rend <= s or rs >= e:
                continue
            a, b = max(s, rs) - rs, min(e, rend) - rs
            r.text = r.text[:a] + (new if first else "") + r.text[b:]
            first = False
    return len(matches)


def all_paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def set_run_font(run, size=None, bold=None, italic=None, small_caps=None):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if small_caps is not None:
        run.font.small_caps = small_caps
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_paragraph(p, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.0, before=0, after=0, **run_kw):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Inches(indent) if indent else Inches(0)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    set_mark_size(p, size)
    for r in p.runs:
        set_run_font(r, size=size, **run_kw)


def section_properties(cols, continuous):
    sp = OxmlElement("w:sectPr")
    if continuous:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), "continuous")
        sp.append(t)
    size = OxmlElement("w:pgSz")
    size.set(qn("w:w"), "11906")
    size.set(qn("w:h"), "16838")
    sp.append(size)
    mar = OxmlElement("w:pgMar")
    for k, v in (("top", 1077), ("right", 812), ("bottom", 2438), ("left", 812),
                 ("header", 708), ("footer", 708), ("gutter", 0)):
        mar.set(qn(f"w:{k}"), str(v))
    sp.append(mar)
    c = OxmlElement("w:cols")
    c.set(qn("w:num"), str(cols))
    c.set(qn("w:space"), "239")
    sp.append(c)
    return sp


def set_mark_size(p, size):
    """Size of the paragraph mark, so empty paragraphs take only `size` points of height."""
    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(size * 2)))


def end_section_at(p, cols, continuous=True):
    """End a section with `cols` columns right after paragraph p.

    The break goes in its own tiny empty paragraph: in a paragraph with text, Word stretches the
    justified last line across the column."""
    holder = p if not p.text.strip() and not p._p.findall(".//" + qn("w:drawing")) else None
    if holder is None:
        new = OxmlElement("w:p")
        p._p.addnext(new)
        holder = docx.text.paragraph.Paragraph(new, p._parent)
    pf = holder.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0)
    set_mark_size(holder, 1)
    holder._p.get_or_add_pPr().append(section_properties(cols, continuous))


def paragraph_after(element, d):
    """The paragraph directly after a body element, inserting an empty one if needed."""
    nxt = element.getnext()
    if nxt is not None and nxt.tag == qn("w:p") and not docx.text.paragraph.Paragraph(nxt, d).text.strip():
        return docx.text.paragraph.Paragraph(nxt, d)
    new = OxmlElement("w:p")
    element.addnext(new)
    return docx.text.paragraph.Paragraph(new, d)


def paragraph_before(element, d):
    prev = element.getprevious()
    while prev is not None and prev.tag != qn("w:p"):
        prev = prev.getprevious()
    return docx.text.paragraph.Paragraph(prev, d)


def column_figures(tmp):
    """Redraw charts for the IEEE layout: column-size canvases, no in-chart titles (captions say it)."""
    figs = importlib.import_module("07_figures")
    figs.FIGURES_DIR = tmp
    figs.title = lambda ax, main, sub: None
    original = plt.subplots

    def smaller(*args, figsize=None, **kw):
        if figsize is not None:
            w, h = CHART_SCALE[current[0]]
            figsize = (figsize[0] * w, figsize[1] * h)
        return original(*args, figsize=figsize, **kw)

    suptitle = matplotlib.figure.Figure.suptitle
    plt.subplots = smaller
    matplotlib.figure.Figure.suptitle = lambda self, *a, **k: None
    names = {c[0] for c in figs.CONDITIONS + figs.TABLE_ONLY} | {"e2_body_only"}
    results = {n: figs.load(n) for n in names}
    for fn in REDRAWN:
        current[0] = fn
        if fn == "fig2_tokens":
            # The caption explains the orange bars; on the shorter canvas the legend would overlap.
            legend = matplotlib.figure.Figure.legend
            matplotlib.figure.Figure.legend = lambda self, *a, **k: None
            figs.fig2_tokens(figs.load("e1_baseline_tokens"))
            matplotlib.figure.Figure.legend = legend
        elif fn == "fig3_confusion":
            figs.fig3_confusion([(h, results[n]) for h, n in (("E1  Full text", "e1_baseline"),
                                 ("E3  Source markers stripped", "e3_stripped"),
                                 ("E3b  + page furniture stripped", "e3b_furniture_stripped"))])
        else:
            getattr(figs, fn)(results)
    plt.subplots = original
    matplotlib.figure.Figure.suptitle = suptitle
    return {fn: tmp / CHART_FILE.get(fn, f"{fn}.png") for fn in REDRAWN}


# ---------------------------------------------------------------- conversion

def convert(tmp):
    d = docx.Document(SRC)
    body = d.element.body
    paras = d.paragraphs

    # 1. Front matter (cover page, contents) out; IEEE title and author block in.
    abstract_h = next(p for p in paras if p.text.strip() == "Abstract" and p.style.name == "Heading 1")
    for el in list(body.iterchildren()):
        if el is abstract_h._p:
            break
        body.remove(el)
    title = abstract_h.insert_paragraph_before(TITLE)
    style_paragraph(title, size=24, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    author_ps = []
    for i, line in enumerate(AUTHOR):
        p = abstract_h.insert_paragraph_before(line)
        style_paragraph(p, size=11 if i == 0 else 10, align=WD_ALIGN_PARAGRAPH.CENTER,
                        italic=(i in (1, 2)), after=(14 if i == len(AUTHOR) - 1 else 0))
        author_ps.append(p)
    end_section_at(author_ps[-1], cols=1, continuous=False)

    # 2. Abstract and index terms.
    abstract_p = docx.text.paragraph.Paragraph(abstract_h._p.getnext(), d)
    body.remove(abstract_h._p)
    lead = abstract_p.runs[0]._r
    new = copy.deepcopy(lead)
    lead.addprevious(new)
    run = docx.text.run.Run(new, abstract_p)
    run.text = "Abstract—"
    style_paragraph(abstract_p, size=9, bold=True, after=6)
    set_run_font(run, size=9, bold=True, italic=True)
    keywords = docx.text.paragraph.Paragraph(abstract_p._p.getnext(), d)
    terms = keywords.text.split(":", 1)[1].strip()
    for r in keywords.runs[1:]:
        r._r.getparent().remove(r._r)
    keywords.runs[0].text = "Index Terms—"
    keywords.add_run(terms[0].lower() + terms[1:])
    style_paragraph(keywords, size=9, bold=True, after=8)
    set_run_font(keywords.runs[0], size=9, bold=True, italic=True)

    # 3. Captions: tables above as "TABLE I" + title in small caps; figures below as "Fig. n."
    table_no = 0
    for p in list(d.paragraphs):
        t = p.text.strip()
        m = re.match(r"Table (\d+)\.\s*(.*)", t)
        if m and p.style.name in ("CaptionText", "Caption"):
            table_no += 1
            caption = m.group(2).rstrip(".")
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = f"TABLE {ROMAN[int(m.group(1))]}"
            p.runs[0].add_break()
            p.add_run(caption)
            style_paragraph(p, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4,
                            bold=False, italic=False, small_caps=True)
            p.paragraph_format.keep_with_next = True
        elif re.match(r"Figure \d+\.", t) and p.style.name in ("CaptionText", "Caption"):
            replace_in_paragraph(p, r"^Figure (\d+)\.", r"Fig. \1.")
            style_paragraph(p, size=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=3, after=8,
                            bold=False, italic=False)

    # 4. Citations and cross-references in body text and tables.
    for p in all_paragraphs(d):
        for pattern, repl in CITATIONS:
            replace_in_paragraph(p, pattern, repl)
        replace_in_paragraph(p, r"Section (\d+)\.(\d+)", lambda m: f"Section {ROMAN[int(m[1])]}-{LETTER[int(m[2])]}")
        replace_in_paragraph(p, r"Sections (\d+) and (\d+)", lambda m: f"Sections {ROMAN[int(m[1])]} and {ROMAN[int(m[2])]}")
        replace_in_paragraph(p, r"Section (\d+)\b", lambda m: f"Section {ROMAN[int(m[1])]}")
        replace_in_paragraph(p, r"\bFigure (\d+)", r"Fig. \1")
        replace_in_paragraph(p, r"\bTable (\d+)\b", lambda m: f"Table {ROMAN[int(m[1])]}")

    # 5. Headings: "6. Results" -> "VI. Results" (centred small caps); "6.2 Probes" -> "B. Probes" (italic).
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            replace_in_paragraph(p, r"^(\d+)\.\s+", lambda m: f"{ROMAN[int(m[1])]}. ")
            style_paragraph(p, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=5,
                            bold=False, italic=False, small_caps=True)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.page_break_before = False
        elif p.style.name == "Heading 2":
            replace_in_paragraph(p, r"^(\d+)\.(\d+)\s+", lambda m: f"{LETTER[int(m[2])]}. ")
            style_paragraph(p, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3,
                            bold=False, italic=True, small_caps=False)
            p.paragraph_format.keep_with_next = True

    # 6. Body text: Times 10 pt, justified, first-line indent; lists without indent.
    skip = {title._p, abstract_p._p, keywords._p} | {p._p for p in author_ps}
    for p in d.paragraphs:
        if p._p in skip or p.style.name.startswith(("Heading", "Caption")):
            continue
        has_image = bool(p._p.findall(".//" + qn("w:drawing")))
        if has_image:
            style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2)
            p.paragraph_format.keep_with_next = True
        elif p.style.name.startswith("List") and re.match(r"(RQ|O)\d+\.", p.text.strip()):
            # Labelled items (RQ1., O1.) carry their own label: drop the list number or bullet.
            p.style = d.styles["Normal"]
            numpr = p._p.pPr.find(qn("w:numPr"))
            if numpr is not None:
                p._p.pPr.remove(numpr)
            style_paragraph(p, size=10, indent=0.17)
        elif p.style.name.startswith("List"):
            style_paragraph(p, size=10)
        else:
            style_paragraph(p, size=10, indent=0.17)

    # 7. Tables: 8 pt, centred; full-width ones span both columns, others fit one.
    for ti, t in enumerate(d.tables):
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        target = FULL_W if ti in FULL_WIDTH_TABLES else COL_W
        widths = [c.width or Inches(1) for c in t.rows[0].cells]
        scale = target / sum(widths)
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = int(widths[ci] * scale)
                for p in cell.paragraphs:
                    style_paragraph(p, size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
                    p.paragraph_format.keep_with_next = False

    # 8. Figures: column-size redraws swapped in; wide ones at full width.
    redrawn = column_figures(tmp)
    for shape, (src, full, redraw) in zip(d.inline_shapes, IMAGES):
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        part = d.part.related_parts[blip.embed]
        path = redrawn[redraw] if redraw else ROOT / src
        data = path.read_bytes()
        part._blob = data
        w_px, h_px = Image.open(io.BytesIO(data)).size
        width = FULL_W if full else COL_W
        shape.width, shape.height = int(width), int(width * h_px / w_px)

    # 9. Full-width objects get their own one-column sections.
    for ti, t in enumerate(d.tables):
        if ti in FULL_WIDTH_TABLES:
            caption = paragraph_before(t._tbl, d)
            end_section_at(paragraph_before(caption._p, d), cols=2)
            end_section_at(paragraph_after(t._tbl, d), cols=1)
    for shape, (_, full, _) in zip(list(d.inline_shapes), IMAGES):
        if full:
            img_p = shape._inline.getparent()
            while img_p.tag != qn("w:p"):
                img_p = img_p.getparent()
            caption = docx.text.paragraph.Paragraph(img_p.getnext(), d)
            end_section_at(paragraph_before(img_p, d), cols=2)
            end_section_at(caption, cols=1)

    # 10. References in IEEE style, in order of first citation.
    ref_h = next(p for p in d.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "References")
    for el in list(ref_h._p.itersiblings()):
        if el.tag == qn("w:p"):
            body.remove(el)
    anchor = ref_h
    for n, segments in enumerate(REFERENCES, 1):
        p = d.add_paragraph()
        anchor._p.addnext(p._p)
        anchor = p
        set_run_font(p.add_run(f"[{n}]\t"), size=8)
        for text, italic in segments:
            set_run_font(p.add_run(text), size=8, italic=italic)
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(0.3), Inches(-0.3)
        pf.tab_stops.add_tab_stop(Inches(0.3))
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(1)

    # 11. Final sections: a continuous break after the last reference balances the two columns on
    # the last page; no page-number footer (IEEE camera-ready style).
    end_section_at(anchor, cols=2)
    final = body.find(qn("w:sectPr"))
    body.remove(final)
    body.append(section_properties(1, True))

    d.core_properties.title = TITLE
    d.core_properties.author = AUTHOR[0]
    d.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    matplotlib.use("Agg")
    with tempfile.TemporaryDirectory() as tmp:
        convert(Path(tmp))
